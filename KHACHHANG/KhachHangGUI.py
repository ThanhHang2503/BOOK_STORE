import tkinter as tk
from tkinter import messagebox, ttk

import pyodbc

from KHACHHANG.KhachHang import KhachHang


class KhachHangGUI(tk.Frame):
    def __init__(self, master=None):
        super().__init__(master)
        self.frame = self  # Thêm thuộc tính frame
        self.customers = []  # Danh sách khách hàng được tải từ SQL
        self.conn = self.ket_noi_sql()  # Thiết lập kết nối SQL
        self.create_widgets()

    def ket_noi_sql(self):
        """ Kết nối với SQL Server """
        try:
            conn = pyodbc.connect(
                "DRIVER={ODBC Driver 17 for SQL Server};"
                "SERVER=DESKTOP-NRE55H1;"
                "DATABASE=DOANPYTHON;"
                "Trusted_Connection=yes;"
            )

            return conn
        except Exception as e:
            messagebox.showerror("Lỗi kết nối SQL", str(e))
            return None

    def load_customers_from_sql(self):
        """ Tải danh sách khách hàng từ SQL Server và cập nhật self.customers """
        if not self.conn:
            messagebox.showerror("Lỗi", "Không thể kết nối SQL!")
            return

        try:
            cursor = self.conn.cursor()
            cursor.execute("SELECT maKH, tenKH, diaChi, dienThoai FROM KhachHang")
            rows = cursor.fetchall()
            self.customers.clear()
            for row in rows:
                customer = {
                    "maKH": row[0],
                    "tenKH": row[1],
                    "diaChi": row[2],
                    "dienThoai": row[3]
                }
                self.customers.append(customer)
        except Exception as e:
            messagebox.showerror("Lỗi SQL", str(e))

    def create_widgets(self):
        # Tạo thanh tiêu đề
        top_bar = tk.Frame(self, bg="#F0F0F0")
        top_bar.grid(row=0, column=0, sticky="ew", pady=5)
        self.grid_columnconfigure(0, weight=1)

        label_title = tk.Label(top_bar, text="Quản Lý Khách Hàng", font=("Arial", 14, "bold"), bg="#F0F0F0")
        label_title.grid(row=0, column=0, padx=10, sticky="w")
        top_bar.grid_columnconfigure(0, weight=1)

        # Tạo các frame cho từng chức năng (bỏ frame_xoa)
        self.frame_them = tk.Frame(self)
        self.frame_sua = tk.Frame(self)
        self.frame_tim = tk.Frame(self)
        self.frame_danhsach = tk.Frame(self)

        # Xây dựng giao diện cho từng chức năng
        self.create_frame_them()
        self.create_frame_sua()
        self.create_frame_tim()
        self.create_frame_danhsach()

    def close_frame(self):
        """Đóng khung Quản Lý Khách Hàng."""
        self.destroy()

    def create_frame_them(self):
        frame = self.frame_them
        tk.Label(frame, text="Mã khách hàng:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
        self.entry_id = tk.Entry(frame)
        self.entry_id.grid(row=0, column=1, padx=10, pady=10)

        tk.Label(frame, text="Tên khách hàng:").grid(row=1, column=0, padx=10, pady=10, sticky="w")
        self.entry_name = tk.Entry(frame)
        self.entry_name.grid(row=1, column=1, padx=10, pady=10)

        tk.Label(frame, text="Địa chỉ:").grid(row=2, column=0, padx=10, pady=10, sticky="w")
        self.entry_address = tk.Entry(frame)
        self.entry_address.grid(row=2, column=1, padx=10, pady=10)

        tk.Label(frame, text="Điện thoại:").grid(row=3, column=0, padx=10, pady=10, sticky="w")
        self.entry_phone = tk.Entry(frame)
        self.entry_phone.grid(row=3, column=1, padx=10, pady=10)

        btn_add = tk.Button(frame, text="Thêm", command=self.add_customer)
        btn_add.grid(row=4, column=0, columnspan=2, padx=10, pady=20)

    def create_frame_sua(self):
        frame = self.frame_sua
        tk.Label(frame, text="Mã khách hàng cần sửa:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
        self.entry_update_id = tk.Entry(frame)
        self.entry_update_id.grid(row=0, column=1, padx=10, pady=10)
        self.entry_update_id.config(state="readonly")
        self.entry_update_id.bind("<FocusOut>", self.prefill_customer_info)

        tk.Label(frame, text="Tên mới:").grid(row=1, column=0, padx=10, pady=10, sticky="w")
        self.entry_update_name = tk.Entry(frame)
        self.entry_update_name.grid(row=1, column=1, padx=10, pady=10)

        tk.Label(frame, text="Địa chỉ mới:").grid(row=2, column=0, padx=10, pady=10, sticky="w")
        self.entry_update_address = tk.Entry(frame)
        self.entry_update_address.grid(row=2, column=1, padx=10, pady=10)

        tk.Label(frame, text="Điện thoại mới:").grid(row=3, column=0, padx=10, pady=10, sticky="w")
        self.entry_update_phone = tk.Entry(frame)
        self.entry_update_phone.grid(row=3, column=1, padx=10, pady=10)

        btn_update = tk.Button(frame, text="Sửa", command=self.update_customer)
        btn_update.grid(row=4, column=0, columnspan=2, padx=10, pady=20)

    def create_frame_tim(self):
        frame = self.frame_tim
        tk.Label(frame, text="Nhập mã khách hàng cần tìm kiếm:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
        self.entry_search = tk.Entry(frame)
        self.entry_search.grid(row=0, column=1, padx=10, pady=10)

        btn_search = tk.Button(frame, text="Tìm", command=self.search_customer)
        btn_search.grid(row=0, column=2, padx=10, pady=10)

        columns = ("Mã KH", "Tên KH", "Địa Chỉ", "Điện Thoại")
        self.tree = ttk.Treeview(frame, columns=columns, show="headings")
        for col in columns:
            self.tree.heading(col, text=col)
        self.tree.grid(row=1, column=0, columnspan=3, padx=10, pady=10, sticky="nsew")

        frame.grid_rowconfigure(1, weight=1)
        frame.grid_columnconfigure(1, weight=1)

    def create_frame_danhsach(self):
        frame = self.frame_danhsach
        columns = ("Mã KH", "Tên KH", "Địa Chỉ", "Điện Thoại")
        self.tree_ds = ttk.Treeview(frame, columns=columns, show="headings")
        for col in columns:
            self.tree_ds.heading(col, text=col)
        self.tree_ds.grid(row=0, column=0, sticky="nsew", padx=10, pady=10)
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(0, weight=1)

        self.tree_ds.bind("<<TreeviewSelect>>", self.prefill_customer_info_from_list)

        # Thêm nút Xuất Excel vào cuối danh sách
        btn_xuat_excel = tk.Button(frame, text="Xuất Excel", command=self.xuatExcel, bg="#2196F3", fg="white")
        btn_xuat_excel.grid(row=1, column=0, sticky="e", padx=10, pady=5)

        self.load_all_customers()

    def load_all_customers(self):
        # Cập nhật dữ liệu từ SQL trước
        self.load_customers_from_sql()

        # Xóa toàn bộ dữ liệu hiện có trong treeview
        for item in self.tree_ds.get_children():
            self.tree_ds.delete(item)
        # Chèn lại dữ liệu từ self.customers
        for customer in self.customers:
            self.tree_ds.insert("", "end", values=(
                customer["maKH"],
                customer["tenKH"],
                customer["diaChi"],
                customer["dienThoai"]
            ))

    # ==========================
    # Các hàm xử lý sự kiện
    # ==========================
    def add_customer(self):
        maKH = self.entry_id.get().strip()
        tenKH = self.entry_name.get().strip()
        diaChi = self.entry_address.get().strip()
        dienThoai = self.entry_phone.get().strip()

        if not maKH or not tenKH:
            messagebox.showerror("Lỗi", "Mã và tên khách hàng không được để trống!")
            return

        # Kiểm tra tồn tại trên SQL (nếu cần) hoặc trong danh sách cục bộ
        if any(cust["maKH"]==maKH for cust in self.customers):
            messagebox.showerror("Lỗi", "Mã khách hàng đã tồn tại!")
            return

        # Thêm khách hàng vào SQL
        if self.conn:
            try:
                cursor = self.conn.cursor()
                cursor.execute(
                    "INSERT INTO KhachHang (maKH, tenKH, diaChi, dienThoai) VALUES (?, ?, ?, ?)",
                    (maKH, tenKH, diaChi, dienThoai)
                )
                self.conn.commit()
            except Exception as e:
                messagebox.showerror("Lỗi SQL", str(e))
                return

        new_customer = {"maKH": maKH, "tenKH": tenKH, "diaChi": diaChi, "dienThoai": dienThoai}
        self.customers.append(new_customer)
        messagebox.showinfo("Thành công", f"Đã thêm khách hàng {maKH}")
        self.load_all_customers()

        # Xóa các trường nhập
        self.entry_id.delete(0, tk.END)
        self.entry_name.delete(0, tk.END)
        self.entry_address.delete(0, tk.END)
        self.entry_phone.delete(0, tk.END)

    def update_customer(self):
        maKH_str = self.entry_update_id.get().strip()
        tenKH = self.entry_update_name.get().strip()
        diaChi = self.entry_update_address.get().strip()
        dienThoai = self.entry_update_phone.get().strip()

        if not maKH_str:
            messagebox.showerror("Lỗi", "Vui lòng nhập mã khách hàng!")
            return

        maKH = maKH_str  # Không cho phép sửa mã khách hàng, lấy đúng mã đang readonly

        # Tải lại danh sách trước khi tìm
        self.load_customers_from_sql()

        # Tìm khách hàng trong self.customers (so sánh dưới dạng chuỗi)
        customer = next((c for c in self.customers if str(c["maKH"]) == str(maKH)), None)
        if not customer:
            messagebox.showerror("Lỗi", f"Không tìm thấy khách hàng với mã {maKH_str}")
            return

        # Kiểm tra xem người dùng đã nhập thông tin mới chưa
        if not tenKH and not diaChi and not dienThoai:
            messagebox.showwarning("Cảnh báo", "Không có thông tin nào được cập nhật!")
            return

        # Hiển thị hộp thoại xác nhận trước khi cập nhật
        if not messagebox.askyesno("Xác nhận", f"Bạn có chắc muốn cập nhật thông tin khách hàng {maKH_str}?"):
            return

        # Cập nhật trong SQL
        try:
            cursor = self.conn.cursor()
            cursor.execute(
                "UPDATE KhachHang SET tenKH = ?, diaChi = ?, dienThoai = ? WHERE maKH = ?",
                (tenKH or customer["tenKH"],
                 diaChi or customer["diaChi"],
                 dienThoai or customer["dienThoai"],
                 maKH)
            )
            self.conn.commit()
        except Exception as e:
            messagebox.showerror("Lỗi SQL", str(e))
            return

        # Cập nhật trong danh sách cục bộ
        customer["tenKH"] = tenKH or customer["tenKH"]
        customer["diaChi"] = diaChi or customer["diaChi"]
        customer["dienThoai"] = dienThoai or customer["dienThoai"]

        messagebox.showinfo("Thành công", f"Đã cập nhật khách hàng {maKH_str}")
        self.load_all_customers()

        # Chọn lại dòng vừa sửa nếu còn tồn tại (so sánh dưới dạng chuỗi)
        for item in self.tree_ds.get_children():
            if str(self.tree_ds.item(item, "values")[0]) == str(maKH_str):
                self.tree_ds.selection_set(item)
                self.tree_ds.see(item)
                break

        # Xóa nội dung các ô nhập sau khi cập nhật thành công
        self.entry_update_id.config(state="normal")
        self.entry_update_id.delete(0, tk.END)
        self.entry_update_id.config(state="readonly")
        self.entry_update_name.delete(0, tk.END)
        self.entry_update_address.delete(0, tk.END)
        self.entry_update_phone.delete(0, tk.END)

    def search_customer(self):
        # Tải lại danh sách trước khi tìm
        self.load_customers_from_sql()

        search_id = self.entry_search.get().strip()
        results = []

        try:
            # Thử chuyển đổi sang số nguyên nếu có thể
            search_id_int = int(search_id)
        except ValueError:
            # Nếu không thể chuyển đổi, giữ nguyên dạng chuỗi
            search_id_int = search_id

        # Tìm khách hàng theo mã KH chính xác
        for cust in self.customers:
            if cust["maKH"]==search_id_int:
                results.append(cust)
                break  # Dừng ngay khi tìm thấy vì mã KH là duy nhất

        # Xóa dữ liệu cũ
        for item in self.tree.get_children():
            self.tree.delete(item)

        # Chèn kết quả
        for customer in results:
            self.tree.insert("", "end", values=(
                customer["maKH"],
                customer["tenKH"],
                customer["diaChi"],
                customer["dienThoai"]
            ))

        if not results:
            messagebox.showinfo("Thông báo", "Không tìm thấy khách hàng với mã này.")

    def prefill_customer_info(self, event=None):
        maKH = self.entry_update_id.get().strip()
        if not maKH:
            return

        self.load_customers_from_sql()

        customer = next((c for c in self.customers if c["maKH"]==maKH), None)
        if customer:
            self.entry_update_name.delete(0, tk.END)
            self.entry_update_address.delete(0, tk.END)
            self.entry_update_phone.delete(0, tk.END)

            self.entry_update_name.insert(0, customer["tenKH"])
            self.entry_update_address.insert(0, customer["diaChi"])
            self.entry_update_phone.insert(0, customer["dienThoai"])
        else:
            self.entry_update_name.delete(0, tk.END)
            self.entry_update_address.delete(0, tk.END)
            self.entry_update_phone.delete(0, tk.END)

    def prefill_customer_info_from_list(self, event=None):
        selected = self.tree_ds.selection()
        if not selected:
            return
        values = self.tree_ds.item(selected[0], "values")
        if values:
            self.entry_update_id.config(state="normal")
            self.entry_update_id.delete(0, tk.END)
            self.entry_update_id.insert(0, values[0])
            self.entry_update_id.config(state="readonly")
            self.entry_update_name.delete(0, tk.END)
            self.entry_update_name.insert(0, values[1])
            self.entry_update_address.delete(0, tk.END)
            self.entry_update_address.insert(0, values[2])
            self.entry_update_phone.delete(0, tk.END)
            self.entry_update_phone.insert(0, values[3])

    def anGiaoDien(self):
        """Ẩn giao diện khách hàng"""
        self.grid_forget()

    def hienThi(self, action=None):
        """Hiển thị giao diện khách hàng với frame tương ứng"""
        self.grid(row=0, column=0, sticky="nsew", padx=100, pady=100)
        
        # Ẩn tất cả các frame
        self.frame_them.grid_forget()
        self.frame_sua.grid_forget()
        self.frame_tim.grid_forget()
        self.frame_danhsach.grid_forget()
        
        # Hiển thị frame tương ứng với action
        if action=="Thêm khách hàng":
            self.frame_them.grid(row=1, column=0, sticky="nsew")
        elif action=="Sửa thông tin khách hàng":
            self.frame_sua.grid(row=1, column=0, sticky="ew", padx=10, pady=10)
            self.frame_danhsach.grid(row=2, column=0, sticky="nsew", padx=10, pady=10)
            self.grid_rowconfigure(2, weight=1)
            self.grid_columnconfigure(0, weight=1)
            self.load_all_customers()
        elif action=="Tìm kiếm khách hàng":
            self.frame_tim.grid(row=1, column=0, sticky="nsew")
        elif action=="Hiển thị danh sách":
            self.frame_danhsach.grid(row=1, column=0, sticky="nsew")
            self.load_all_customers()

    def xuatExcel(self):
        """Xuất danh sách khách hàng ra file Excel"""
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter

            # Lấy danh sách khách hàng
            self.load_customers_from_sql()
            ds = self.customers
            if not ds:
                messagebox.showwarning("Cảnh báo", "Không có dữ liệu khách hàng để xuất!")
                return

            from tkinter import filedialog
            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                title="Lưu file Excel"
            )
            if not file_path:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Danh sách khách hàng"

            headers = ["Mã KH", "Tên KH", "Địa chỉ", "Điện thoại"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for row, kh in enumerate(ds, 2):
                ws.cell(row=row, column=1, value=kh["maKH"])
                ws.cell(row=row, column=2, value=kh["tenKH"])
                ws.cell(row=row, column=3, value=kh["diaChi"])
                ws.cell(row=row, column=4, value=kh["dienThoai"])

            for col in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(col)].width = 18

            wb.save(file_path)
            messagebox.showinfo("Thành công", f"Đã xuất danh sách khách hàng thành công!\nFile được lưu tại: {file_path}")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Có lỗi xảy ra khi xuất file Excel:\n{str(e)}")


if __name__=="__main__":
    root = tk.Tk()
    root.title("Quản lý khách hàng")
    root.geometry("800x600")
    app = KhachHangGUI(root)
    root.mainloop()